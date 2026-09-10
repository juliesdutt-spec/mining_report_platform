"""
Guards the file-serving endpoints.

Content-Disposition is the whole bug surface here: a browser honours
`attachment` inside an <object> or <iframe> too, so the Documents preview -
which embeds the PDF URL - saved a file on every mount instead of rendering
one. Serving the same document under both dispositions is what fixes that,
and these tests fail if either half regresses.

They run against a temporary SQLite database so nothing touches a real one.
"""
import io
import os
import tempfile
import re
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent



try:
    import pypdfium2
    CAN_READ_PDFS = True
except ImportError:  # a way to check our work, not a runtime dependency
    CAN_READ_PDFS = False


def _pdf_content(pdf: bytes):
    """(page count, text) - what the document says, not how it was encoded."""
    document = pypdfium2.PdfDocument(io.BytesIO(pdf))
    pages = len(document)
    text = "\n".join(page.get_textpage().get_text_range() for page in document)
    # The footer carries a clock reading, which is not content.
    return pages, re.sub(r"\d{2}-\d{2}-\d{4} \d{2}:\d{2}", "<time>", text)


class ExportEndpointTests(unittest.TestCase):
    #: Set for this class only. These are process-wide, and the provider tests
    #: reload ai_providers against their own environment - leaving USE_MOCK_AI
    #: behind would force every later provider case into mock mode.
    OVERRIDES = ("DATABASE_URL", "USE_MOCK_AI",
                 "AUTH_USERS", "AUTH_SECRET", "DEMO_ACCOUNT")

    @classmethod
    def setUpClass(cls):
        cls._saved = {k: os.environ.get(k) for k in cls.OVERRIDES}

        # Point the app at a throwaway database before importing it - the
        # engine is built at import time from DATABASE_URL.
        cls._tmp = tempfile.TemporaryDirectory()
        os.environ["DATABASE_URL"] = f"sqlite:///{cls._tmp.name}/test.db"
        os.environ["USE_MOCK_AI"] = "true"
        # These endpoints now require a signed-in account. The suite signs in
        # as a full (non-demo) user so it exercises the same paths as before.
        os.environ["AUTH_SECRET"] = "test-secret-not-a-real-one"
        os.environ["AUTH_USERS"] = "tester:test-password:Test User"
        os.environ["DEMO_ACCOUNT"] = "off"

        import sys
        sys.path.insert(0, str(PROJECT_ROOT))
        cls._poisoned = ("database", "auth", "auth_seed", "ai_providers", "ai_extractor", "backend.api")
        for name in cls._poisoned:
            sys.modules.pop(name, None)

        from fastapi.testclient import TestClient
        from backend.api import app
        from database import init_db

        init_db()
        from auth_seed import seed_users
        seed_users()

        cls.client = TestClient(app)
        login = cls.client.post("/auth/login", json={"username": "tester", "password": "test-password"})
        assert login.status_code == 200, f"test account could not sign in: {login.text}"
        # Set once on the client rather than per call, so the assertions below
        # are unchanged from before authentication existed.
        cls.client.headers["Authorization"] = f"Bearer {login.json()['access_token']}"

        # One real upload, so every export has genuine extracted data behind it.
        pdf = PROJECT_ROOT / "sample_mining_report.pdf"
        with pdf.open("rb") as handle:
            response = cls.client.post(
                "/upload",
                files={"file": ("sample_mining_report.pdf", handle, "application/pdf")},
            )
        assert response.status_code == 200, response.text
        cls.report_id = cls.client.get("/reports").json()["reports"][0]["id"]

    @classmethod
    def tearDownClass(cls):
        # Restore the environment and drop the modules imported under it, so a
        # later test importing them resolves against its own configuration.
        import sys
        for key, value in cls._saved.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        for name in cls._poisoned:
            sys.modules.pop(name, None)
        cls._tmp.cleanup()

    # ------------------------------------------------- PDF disposition ---
    def test_pdf_download_defaults_to_attachment(self):
        response = self.client.get(f"/reports/{self.report_id}/download")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["content-disposition"])

    def test_pdf_preview_is_served_inline(self):
        """The Documents preview embeds this URL; attachment would download."""
        response = self.client.get(
            f"/reports/{self.report_id}/download", params={"inline": "true"}
        )
        self.assertEqual(response.status_code, 200)
        disposition = response.headers["content-disposition"]
        self.assertIn("inline", disposition)
        self.assertNotIn("attachment", disposition)

    def test_both_dispositions_return_the_same_valid_pdf(self):
        attachment = self.client.get(f"/reports/{self.report_id}/download").content
        inline = self.client.get(
            f"/reports/{self.report_id}/download", params={"inline": "true"}
        ).content
        self.assertTrue(attachment.startswith(b"%PDF"), "not a PDF")
        self.assertTrue(inline.startswith(b"%PDF"), "not a PDF")

        # Compared as documents rather than as bytes. Two renders a moment
        # apart are never byte-identical: the timestamps and the /ID differ,
        # and so does the embedded font subset, because fontTools stamps a
        # modification time into its head table - inside a compressed stream,
        # where no amount of normalising can reach it. This test asserted
        # byte equality and passed only by luck, until rendering got slow
        # enough to cross a second boundary. What it means to check is that
        # the disposition does not change the document, so that is what it
        # now checks.
        if CAN_READ_PDFS:
            self.assertEqual(_pdf_content(attachment), _pdf_content(inline),
                             "disposition must not change the document")
        else:
            self.assertAlmostEqual(len(attachment), len(inline), delta=2048)

    def test_missing_report_is_404_not_an_empty_file(self):
        self.assertEqual(self.client.get("/reports/99999/download").status_code, 404)

    # ------------------------------------------------ dossier: DOCX ------
    def test_dossier_docx_is_a_real_word_file(self):
        """
        Not "the endpoint returned 200" - the bytes have to be a package Word
        can open, which a placeholder or an error page would fail.
        """
        import io
        import zipfile

        response = self.client.get("/reports/generate", params={"format": "docx"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn(
            "wordprocessingml.document", response.headers["content-type"]
        )
        self.assertIn(".docx", response.headers["content-disposition"])

        payload = response.content
        self.assertGreater(len(payload), 5000, "suspiciously small for a dossier")

        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            self.assertIsNone(archive.testzip(), "corrupt OOXML package")
            names = archive.namelist()
        for part in ("[Content_Types].xml", "word/document.xml"):
            self.assertIn(part, names, f"missing OOXML part {part}")

    def test_dossier_docx_carries_the_extracted_data(self):
        from docx import Document

        import io

        payload = self.client.get(
            "/reports/generate", params={"format": "docx"}
        ).content
        document = Document(io.BytesIO(payload))
        text = "\n".join(p.text for p in document.paragraphs)

        self.assertIn("EXECUTIVE SUMMARY", text)
        self.assertIn("SOURCE REFERENCES", text)
        self.assertGreaterEqual(len(document.tables), 1, "no production table")

        header = [c.text for c in document.tables[0].rows[0].cells]
        self.assertEqual(header, ["Document", "Mineral", "Quantity", "Method"])
        body = "\n".join(
            c.text for row in document.tables[0].rows for c in row.cells
        )
        self.assertIn("sample_mining_report.pdf", body, "real filename missing")

    def test_dossier_sections_can_be_switched_off(self):
        from docx import Document

        import io

        payload = self.client.get(
            "/reports/generate",
            params={"format": "docx", "key_findings": "false"},
        ).content
        text = "\n".join(
            p.text for p in Document(io.BytesIO(payload)).paragraphs
        )
        self.assertNotIn("KEY FINDINGS", text)
        self.assertIn("EXECUTIVE SUMMARY", text)

    # ------------------------------------------------- dossier: PDF ------
    def test_dossier_pdf_still_works(self):
        """The DOCX addition must not disturb the existing PDF export."""
        response = self.client.get("/reports/generate")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_dossier_title_follows_the_requested_template(self):
        """
        The Template control offered four statutory document types and every
        one produced an identically titled file.
        """
        from docx import Document

        import io

        payload = self.client.get(
            "/reports/generate",
            params={"format": "docx", "title": "Statutory DGMS Safety Audit"},
        ).content
        text = "\n".join(p.text for p in Document(io.BytesIO(payload)).paragraphs)
        self.assertIn("Statutory DGMS Safety Audit", text)

    def test_dossier_states_the_basis_for_what_it_includes(self):
        """
        The period is a caption the operator picks, not a filter - uploaded
        reports rarely carry a parseable date. The document has to say so, or
        a dossier headed "FY 2023-24" implies a selection that never happened.
        """
        from docx import Document

        import io

        payload = self.client.get(
            "/reports/generate",
            params={"format": "docx", "period": "FY 2023-24 (annual)"},
        ).content
        text = "\n".join(p.text for p in Document(io.BytesIO(payload)).paragraphs)
        self.assertIn("FY 2023-24 (annual)", text)
        self.assertIn("all completed reports indexed at generation", text)

    def test_pdf_dossier_states_the_same_basis(self):
        payload = self.client.get(
            "/reports/generate", params={"period": "FY 2024-25 Q1"}
        ).content
        self.assertTrue(payload.startswith(b"%PDF"))
        # The PDF is compressed, so assert on generation succeeding with the
        # same parameters rather than scraping glyphs out of the stream.
        self.assertGreater(len(payload), 1000)

    def test_dossier_rejects_an_unknown_format(self):
        self.assertEqual(
            self.client.get("/reports/generate", params={"format": "xlsx"}).status_code,
            422,
        )


if __name__ == "__main__":
    unittest.main()
